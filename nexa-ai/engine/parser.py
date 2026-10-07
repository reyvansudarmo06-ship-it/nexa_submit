from pathlib import Path

from PyPDF2 import PdfReader
from docx import Document
from openpyxl import load_workbook


def parse_pdf(file_path):
    reader = PdfReader(file_path)

    text = []

    for page in reader.pages:
        content = page.extract_text()

        if content:
            text.append(content)

    return "\n".join(text).strip()


def parse_docx(file_path):
    document = Document(file_path)

    text = []

    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            text.append(paragraph.text)

    return "\n".join(text).strip()


def parse_xlsx(file_path):
    workbook = load_workbook(
        file_path,
        read_only=True,
        data_only=True
    )

    text = []

    for worksheet in workbook.worksheets:

        text.append(
            f"[Sheet: {worksheet.title}]"
        )

        for row in worksheet.iter_rows(
            values_only=True
        ):
            values = []

            for value in row:
                if value is not None:
                    values.append(str(value))

            if values:
                text.append(
                    " | ".join(values)
                )

    return "\n".join(text).strip()


def parse_file(file_path, mime_type=None):

    path = Path(file_path)

    extension = path.suffix.lower()

    if extension == ".pdf":
        return parse_pdf(file_path)

    if extension == ".docx":
        return parse_docx(file_path)

    if extension in [".xlsx", ".xlsm"]:
        return parse_xlsx(file_path)

    raise ValueError(
        f"Format file belum didukung: {extension}"
    )